"""
Motor del generador de CV. Python puro: no importa Streamlit.

El currículo es un diccionario plano (ver `nuevo()`), con dos listas de
fichas: experiencias y formación. Aquí vive todo lo que no dibuja:

    a_oracion(denominacion)   nombre de catálogo -> nombre para el currículo
    ordena_por_fechas(fichas) más reciente primero
    texto_plano(cv)           vista previa en texto
    documento_docx(cv)        el Word, en bytes
    documento_pdf(cv)         el PDF, en bytes, con la misma maquetación
"""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROJO = RGBColor(0xD1, 0x12, 0x2E)
GRIS = RGBColor(0x55, 0x55, 0x55)


def nuevo():
    return {
        "nombre": "", "telefono": "", "email": "", "localidad": "",
        "permiso": "", "disponibilidad": "",
        "perfil": "",
        "experiencias": [],
        "formacion": [],
        "idiomas": "", "informatica": "", "otros": "",
    }


def experiencia(codigo="", denominacion="", motivo=""):
    """Una ficha de experiencia. `codigo` es el SISPE si viene del codificador."""
    return {
        "codigo": codigo,
        "denominacion": denominacion,
        "motivo": motivo,
        "sector": "",
        "puesto": a_oracion(denominacion) if denominacion else "",
        "contexto": "",
        "funciones": "",
        "desde": "",
        "hasta": "",
    }


def formacion(titulo="", centro="", anio=""):
    return {"titulo": titulo, "centro": centro, "anio": anio}


SUELTAS_ES = ("r", "n", "l", "d", "s", "z", "j")


def a_oracion(denom):
    """Pasa la denominación oficial a algo que se pueda poner en un CV.

    El catálogo va en MAYÚSCULAS y en plural ("CAMAREROS DE PISO"), porque así
    consta oficialmente y así debe verse en el codificador. Pero un currículo
    se escribe en singular y en tipo oración ("Camarero de piso"). El singular
    es aproximado a propósito: el campo queda editable y la persona lo ajusta.
    """
    base = re.sub(r"\s*\([^)]*\)", "", denom or "").strip().lower()
    base = re.sub(r",?\s*(en general|en gral\.?|n\.c\.o\.p\.?)\s*$", "", base).strip()
    base = base.strip(" ,;")
    if not base:
        return ""

    def singular(palabra):
        if palabra.endswith("es") and len(palabra) > 4 and palabra[-3] in SUELTAS_ES:
            return palabra[:-2]
        if palabra.endswith("s") and len(palabra) > 3:
            return palabra[:-1]
        return palabra

    piezas = base.split()
    piezas[0] = "-".join(singular(x) for x in piezas[0].split("-"))
    return " ".join(piezas).capitalize()


def anio(texto):
    """El último año de cuatro cifras que aparezca en el texto, o 0."""
    m = re.findall(r"(?:19|20)\d{2}", texto or "")
    return int(m[-1]) if m else 0


def ordena_por_fechas(fichas):
    """Más reciente primero, que es como se lee un currículo.

    Las que no tienen ningún año se quedan al final y conservan el orden en
    que se añadieron, para poder colocarlas a mano con las flechas.
    """
    fichas.sort(
        key=lambda e: (
            0 if (anio(e.get("hasta")) or anio(e.get("desde"))) else 1,
            -(anio(e.get("hasta")) or anio(e.get("desde"))),
            -anio(e.get("desde")),
        ),
    )


def titulo_experiencia(e):
    return e.get("puesto") or a_oracion(e.get("denominacion", "")) or "Experiencia sin nombre"


def periodo(e):
    if not (e.get("desde") or e.get("hasta")):
        return ""
    rango = " - ".join(x for x in (e.get("desde", ""), e.get("hasta", "")) if x)
    a1, a2 = anio(e.get("desde")), anio(e.get("hasta"))
    if a1 and a2 and a2 >= a1:
        n = max(1, a2 - a1)
        return f"{n} año{'s' if n != 1 else ''} - {rango}"
    return rango


def _sectores_que_agrupan(experiencias):
    # El sector solo es un agrupador: se escribe si reúne dos o más puestos.
    # Con uno solo sería un título para una línea, que gasta espacio sin
    # aportar nada, y en un currículo de una página el espacio es el límite.
    cuenta = {}
    for e in experiencias:
        s = (e.get("sector") or "").strip().upper()
        if s:
            cuenta[s] = cuenta.get(s, 0) + 1
    return {k for k, v in cuenta.items() if v >= 2}


def experiencias_agrupadas(experiencias):
    """[(sector o None, ficha), ...] en orden, con el sector solo en su primera ficha."""
    agrupan = _sectores_que_agrupan(experiencias)
    salida, actual = [], None
    for e in experiencias:
        sector = (e.get("sector") or "").strip().upper()
        if sector in agrupan and sector != actual:
            salida.append((sector, e))
            actual = sector
        else:
            if sector not in agrupan:
                actual = None
            salida.append((None, e))
    return salida


def contacto(cv):
    piezas = [cv.get("telefono"), cv.get("email"), cv.get("localidad")]
    return "  ·  ".join(p.strip() for p in piezas if p and p.strip())


def texto_plano(cv):
    lineas = []
    if cv.get("nombre"):
        lineas += [cv["nombre"].upper(), contacto(cv), ""]
    if cv.get("perfil"):
        lineas += ["PERFIL PROFESIONAL", cv["perfil"], ""]
    if cv.get("experiencias"):
        lineas.append("EXPERIENCIA LABORAL")
        for sector, e in experiencias_agrupadas(cv["experiencias"]):
            if sector:
                lineas.append(f"\n{sector}")
            p = periodo(e)
            lineas.append(f"·   {titulo_experiencia(e)}{f' ({p})' if p else ''}")
            if e.get("contexto"):
                lineas.append(e["contexto"])
            if e.get("funciones"):
                lineas.append(f"Funciones: {e['funciones']}")
        lineas.append("")
    if cv.get("formacion"):
        lineas.append("FORMACIÓN")
        for f in cv["formacion"]:
            detalle = " · ".join(x for x in (f.get("centro"), f.get("anio")) if x)
            lineas.append(f"·   {f.get('titulo', '')}{f' ({detalle})' if detalle else ''}")
        lineas.append("")
    for rotulo, clave in (("IDIOMAS", "idiomas"), ("INFORMÁTICA", "informatica"),
                          ("OTROS DATOS", "otros")):
        if cv.get(clave):
            lineas += [rotulo, cv[clave], ""]
    extra = [x for x in (cv.get("permiso"), cv.get("disponibilidad")) if x]
    if extra and not cv.get("otros"):
        lineas += ["OTROS DATOS", " · ".join(extra), ""]
    return "\n".join(lineas).strip()


def _seccion(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(texto.upper())
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = ROJO
    # Línea bajo el rótulo, sin tablas: un borde inferior de párrafo.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = p._p.get_or_add_pPr()
    borde = OxmlElement("w:pBdr")
    abajo = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "D1122E")):
        abajo.set(qn(k), v)
    borde.append(abajo)
    pPr.append(borde)
    return p


def _parrafo(doc, texto, negrita=False, gris=False, tamano=10, antes=0, despues=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    r = p.add_run(texto)
    r.bold = negrita
    r.font.size = Pt(tamano)
    if gris:
        r.font.color.rgb = GRIS
    return p


def documento_docx(cv):
    """El currículo en Word, listo para descargar. Devuelve bytes."""
    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run((cv.get("nombre") or "Nombre y apellidos").strip())
    r.bold = True
    r.font.size = Pt(20)
    if contacto(cv):
        _parrafo(doc, contacto(cv), gris=True, tamano=9.5, despues=4)

    if cv.get("perfil"):
        _seccion(doc, "Perfil profesional")
        _parrafo(doc, cv["perfil"])

    if cv.get("experiencias"):
        _seccion(doc, "Experiencia laboral")
        for sector, e in experiencias_agrupadas(cv["experiencias"]):
            if sector:
                _parrafo(doc, sector, negrita=True, gris=True, tamano=9, antes=4)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(titulo_experiencia(e))
            r.bold = True
            r.font.size = Pt(10.5)
            if periodo(e):
                r2 = p.add_run(f"   {periodo(e)}")
                r2.font.size = Pt(9.5)
                r2.font.color.rgb = GRIS
            if e.get("contexto"):
                _parrafo(doc, e["contexto"], gris=True, tamano=9.5, despues=0)
            if e.get("funciones"):
                _parrafo(doc, e["funciones"], despues=2)

    if cv.get("formacion"):
        _seccion(doc, "Formación")
        for f in cv["formacion"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f.get("titulo", ""))
            r.bold = True
            detalle = " · ".join(x for x in (f.get("centro"), f.get("anio")) if x)
            if detalle:
                r2 = p.add_run(f"   {detalle}")
                r2.font.size = Pt(9.5)
                r2.font.color.rgb = GRIS

    for rotulo, clave in (("Idiomas", "idiomas"), ("Informática", "informatica")):
        if cv.get(clave):
            _seccion(doc, rotulo)
            _parrafo(doc, cv[clave])

    otros = [x for x in (cv.get("permiso"), cv.get("disponibilidad"), cv.get("otros")) if x]
    if otros:
        _seccion(doc, "Otros datos")
        _parrafo(doc, " · ".join(otros))

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# PDF: la misma maquetación que el Word, con reportlab
# ---------------------------------------------------------------------------

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape as _esc

_ROJO_PDF = colors.HexColor("#D1122E")
_GRIS_PDF = colors.HexColor("#555555")

_E = {
    "nombre": ParagraphStyle("nombre", fontName="Helvetica-Bold", fontSize=20, leading=24,
                             alignment=TA_LEFT, spaceAfter=2),
    "contacto": ParagraphStyle("contacto", fontName="Helvetica", fontSize=9.5, leading=12,
                               textColor=_GRIS_PDF, spaceAfter=6),
    "seccion": ParagraphStyle("seccion", fontName="Helvetica-Bold", fontSize=10.5, leading=13,
                              textColor=_ROJO_PDF, spaceBefore=10, spaceAfter=1),
    "normal": ParagraphStyle("normal", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=2),
    "sector": ParagraphStyle("sector", fontName="Helvetica-Bold", fontSize=9, leading=11,
                             textColor=_GRIS_PDF, spaceBefore=4, spaceAfter=1),
    "puesto": ParagraphStyle("puesto", fontName="Helvetica", fontSize=10.5, leading=13,
                             spaceBefore=3, spaceAfter=0),
    "contexto": ParagraphStyle("contexto", fontName="Helvetica", fontSize=9.5, leading=12,
                               textColor=_GRIS_PDF, spaceAfter=0),
    "titulo_f": ParagraphStyle("titulo_f", fontName="Helvetica", fontSize=10, leading=13, spaceAfter=1),
}


def _p(texto, estilo):
    return Paragraph(_esc(texto or "").replace("\n", "<br/>"), _E[estilo])


def _seccion_pdf(texto):
    return [
        _p(texto.upper(), "seccion"),
        HRFlowable(width="100%", thickness=0.75, color=_ROJO_PDF, spaceBefore=0, spaceAfter=3),
    ]


def documento_pdf(cv):
    """El currículo en PDF. Devuelve bytes. Misma maquetación que el Word."""
    salida = io.BytesIO()
    doc = SimpleDocTemplate(
        salida, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.8 * cm, bottomMargin=1.8 * cm,
        title=(cv.get("nombre") or "Currículo").strip(), author="", subject="Currículo",
    )
    f = [_p((cv.get("nombre") or "Nombre y apellidos").strip(), "nombre")]
    if contacto(cv):
        f.append(_p(contacto(cv), "contacto"))

    if cv.get("perfil"):
        f += _seccion_pdf("Perfil profesional")
        f.append(_p(cv["perfil"], "normal"))

    if cv.get("experiencias"):
        f += _seccion_pdf("Experiencia laboral")
        for sector, e in experiencias_agrupadas(cv["experiencias"]):
            if sector:
                f.append(_p(sector, "sector"))
            linea = f"<b>{_esc(titulo_experiencia(e))}</b>"
            if periodo(e):
                linea += f'   <font size="9.5" color="#555555">{_esc(periodo(e))}</font>'
            f.append(Paragraph(linea, _E["puesto"]))
            if e.get("contexto"):
                f.append(_p(e["contexto"], "contexto"))
            if e.get("funciones"):
                f.append(_p(e["funciones"], "normal"))

    if cv.get("formacion"):
        f += _seccion_pdf("Formación")
        for x in cv["formacion"]:
            linea = f"<b>{_esc(x.get('titulo', ''))}</b>"
            detalle = " · ".join(v for v in (x.get("centro"), x.get("anio")) if v)
            if detalle:
                linea += f'   <font size="9.5" color="#555555">{_esc(detalle)}</font>'
            f.append(Paragraph(linea, _E["titulo_f"]))

    for rotulo, clave in (("Idiomas", "idiomas"), ("Informática", "informatica")):
        if cv.get(clave):
            f += _seccion_pdf(rotulo)
            f.append(_p(cv[clave], "normal"))

    otros = [x for x in (cv.get("permiso"), cv.get("disponibilidad"), cv.get("otros")) if x]
    if otros:
        f += _seccion_pdf("Otros datos")
        f.append(_p(" · ".join(otros), "normal"))

    f.append(Spacer(1, 1))
    doc.build(f)
    return salida.getvalue()
