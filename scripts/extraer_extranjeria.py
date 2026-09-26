"""
Saca del Excel de códigos de extranjería los CSV que lee el codificador.

USO
    ~/.venvs/sispe/bin/python scripts/extraer_extranjeria.py RUTA_AL_XLSX

El Excel («CODIGOS_Autorizaciones_extranjeria_SEPT_2026.xlsx», en la carpeta
de Drive del curso 2026CE310502) es el documento de la oficina y no se
versiona: aquí entran solo los CSV que salen de él, en
`herramientas/extranjeria/datos/`. Cuando el Excel cambie, se vuelve a pasar
este script y se mira el diff antes de hacer commit.

Qué hoja alimenta qué archivo:

    reglas (A:H)          reglas.csv          el árbol de decisión, una fila por opción
    reglas (J:L)          documentos.csv      los siete documentos del paso 1 y dónde mirar
    reglas (N:Q)          duraciones.csv      cuánto suele durar la tarjeta de cada código
    TABLA DE CÓDIGOS      codigos.csv         los códigos con sus siete columnas
    EN TRÁMITE            tramite.csv         qué se inscribe estando en trámite y qué no
    FECHAS                plazos.csv, bajas.csv, duraciones_tarjeta.csv
    TEXTOS TIE            textos_tie.csv      qué pone la tarjeta y qué código suele ser
    ENLACES               enlaces.csv
    NOVEDADES             novedades.csv
    (varias)              notas.json          títulos, leyenda y las notas sueltas

Las hojas IDENTIFICAR y CONSULTAR CÓDIGO son fórmulas: su lógica vive en
`herramientas/extranjeria/motor.py`, no en los datos.

No lo usa la app. Necesita openpyxl (está en ~/.venvs/sispe).
"""

import csv
import json
import os
import sys
from collections import Counter

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DESTINO = os.path.join(RAIZ, "herramientas", "extranjeria", "datos")


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def fila(ws, r, cols):
    return [texto(ws[f"{c}{r}"].value) for c in cols]


def escribe(nombre, cabecera, filas):
    with open(os.path.join(DESTINO, nombre), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(cabecera)
        w.writerows(filas)
    print(f"  {nombre:24} {len(filas):3} filas")
    return filas


def extrae(ruta):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=False)
    os.makedirs(DESTINO, exist_ok=True)
    notas = {}

    # --- reglas -----------------------------------------------------------
    ws = wb["reglas"]
    reglas, documentos, duraciones = [], [], []
    for r in range(2, ws.max_row + 1):
        if texto(ws[f"A{r}"].value):
            reglas.append(fila(ws, r, "ABCDEFGH"))
        if texto(ws[f"J{r}"].value):
            documentos.append(fila(ws, r, "JKL"))
        if texto(ws[f"N{r}"].value):
            duraciones.append(fila(ws, r, "NOPQ"))
    escribe("reglas.csv", ["lista", "opcion", "tipo", "cod1", "cod2", "vigencia", "aviso", "marca"], reglas)
    escribe("documentos.csv", ["documento", "lista", "donde_mirar"], documentos)
    escribe("duraciones.csv", ["codigo", "minimo", "maximo", "normal"], duraciones)

    repetidas = [o for o, n in Counter(x[1] for x in reglas).items() if n > 1]
    if repetidas:
        print("  AVISO: opciones repetidas en `reglas` (el Excel coge la primera):")
        for o in repetidas:
            print(f"    · {o}")

    # --- TABLA DE CÓDIGOS ---------------------------------------------------
    ws = wb["TABLA DE CÓDIGOS"]
    notas["titulo"] = texto(ws["A1"].value)
    notas["fuente"] = texto(ws["A2"].value)
    codigos, grupo = [], ""
    for r in range(5, ws.max_row + 1):
        a, b = texto(ws[f"A{r}"].value), texto(ws[f"B{r}"].value)
        if a and not b:
            grupo = a
        elif a:
            codigos.append([a, grupo] + fila(ws, r, "BCDEFGH"))
    escribe("codigos.csv", ["codigo", "grupo", "que_es", "puede_trabajar", "restriccion",
                            "colectivo", "se_inscribe_con", "fecha_fin", "ojo"], codigos)

    # --- EN TRÁMITE -----------------------------------------------------------
    ws = wb["EN TRÁMITE"]
    tramite, apartado, notas["tramite"] = [], "", []
    for r in range(3, ws.max_row + 1):
        a, b, c = fila(ws, r, "ABC")
        if a.startswith("✔"):
            apartado = "SI"
        elif a.startswith("✘"):
            apartado = "NO"
        elif a.startswith("•"):
            notas["tramite"].append(a[1:].strip())
        elif a:
            tramite.append([apartado, a, b, c])
    escribe("tramite.csv", ["apartado", "situacion", "codigo", "como_se_graba"], tramite)

    # --- FECHAS -----------------------------------------------------------------
    ws = wb["FECHAS"]
    plazos, bajas, dur_tarjeta, seccion = [], [], [], 1
    notas["plazos"], notas["duraciones_tarjeta"] = [], []
    for r in range(4, ws.max_row + 1):
        a, b, c, d = fila(ws, r, "ABCD")
        if a.startswith("¿CUÁNDO"):
            seccion = 2
        elif a.startswith("CAUSAS DE BAJA"):
            seccion = 3
        elif a.startswith("DURACIONES HABITUALES"):
            seccion = 4
        elif a.startswith("•"):
            notas["plazos"].append(a[1:].strip())
        elif seccion == 1 and a and b and d and a != "SUPUESTO":
            plazos.append([a, b, d])
        elif seccion == 1 and a.startswith("Fechas fijas"):
            notas["plazos"].append(a)
        elif seccion == 3 and a and b:
            bajas.append([a, b])
        elif seccion == 4 and a and c and a != "DURACIÓN":
            dur_tarjeta.append([a, c])
        elif seccion == 4 and a and not c:
            notas["duraciones_tarjeta"].append(a)
    escribe("plazos.csv", ["supuesto", "codigo", "regla"], plazos)
    escribe("bajas.csv", ["causa", "codigo"], bajas)
    escribe("duraciones_tarjeta.csv", ["duracion", "codigos"], dur_tarjeta)

    # --- TEXTOS TIE ---------------------------------------------------------------
    ws = wb["TEXTOS TIE"]
    notas["textos_tie"] = texto(ws["A2"].value)
    textos = []
    for r in range(4, ws.max_row + 1):
        f = fila(ws, r, "ABCDE")
        if f[1] and f[0] not in ("DÓNDE",):
            textos.append(f)
    escribe("textos_tie.csv", ["donde", "texto", "codigo", "ojo", "fuente"], textos)

    # --- ENLACES ------------------------------------------------------------------
    ws = wb["ENLACES"]
    enlaces, apartado = [], ""
    for r in range(2, ws.max_row + 1):
        a, b = fila(ws, r, "AB")
        if a and not b:
            apartado = a
        elif a:
            enlaces.append([apartado, a, b])
    escribe("enlaces.csv", ["apartado", "nombre", "dato"], enlaces)

    # --- NOVEDADES ------------------------------------------------------------------
    ws = wb["NOVEDADES"]
    notas["novedades"] = texto(ws["A1"].value)
    novedades = [fila(ws, r, "AB") for r in range(2, ws.max_row + 1)
                 if texto(ws[f"A{r}"].value) and texto(ws[f"B{r}"].value)]
    escribe("novedades.csv", ["codigo", "cambio"], novedades)

    # --- lo suelto ----------------------------------------------------------------------
    ws = wb["CONSULTAR CÓDIGO"]
    notas["leyenda"] = [fila(ws, r, "BC") for r in range(15, 18)]
    notas["aviso"] = texto(wb["IDENTIFICAR"]["B20"].value)
    with open(os.path.join(DESTINO, "notas.json"), "w", encoding="utf-8") as f:
        json.dump(notas, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print("  notas.json")


if __name__ == "__main__":
    if len(sys.argv) != 2 or not os.path.exists(sys.argv[1]):
        sys.exit(__doc__)
    print(f"Extrayendo de {os.path.basename(sys.argv[1])} a {os.path.relpath(DESTINO, RAIZ)}/")
    extrae(sys.argv[1])
